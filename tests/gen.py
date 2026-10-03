"""生成每种支持的输入格式的样本文件（中文文件名+空格）。"""
import os, subprocess, sys
from pathlib import Path
from PIL import Image, ImageDraw
import pillow_heif; pillow_heif.register_heif_opener()
OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=True)
MARK = "MiniBox测试文字"
def ff(*a): subprocess.run(["ffmpeg","-hide_banner","-loglevel","error","-y",*a], check=True)
# ---- 图片
def pic(mode="RGBA", size=(640,480)):
    im = Image.new(mode, size, (30,120,200,255) if mode=="RGBA" else (30,120,200))
    d = ImageDraw.Draw(im); d.ellipse((100,80,540,400), fill=(255,200,0,180) if mode=="RGBA" else (255,200,0))
    return im
pic().save(OUT/"样图 透明.png")
pic("RGB").save(OUT/"样图.jpg", quality=90); pic("RGB").save(OUT/"样图2.jpeg")
pic().save(OUT/"样图.webp"); pic("RGB").save(OUT/"样图.bmp"); pic("RGB").save(OUT/"样图.tif"); pic("RGB").save(OUT/"样图2.tiff")
pic().save(OUT/"样图.ico", sizes=[(64,64)])
pic("RGB").save(OUT/"样图.heic"); pic("RGB").save(OUT/"样图2.heif"); pic("RGB").save(OUT/"样图.avif")
frames = [pic("RGB").rotate(i*30) for i in range(6)]
frames[0].save(OUT/"动图.gif", save_all=True, append_images=frames[1:], duration=120, loop=0)
# ---- 视频（5 秒）
src = ["-f","lavfi","-i","testsrc2=duration=5:size=640x360:rate=25","-f","lavfi","-i","sine=frequency=440:duration=5"]
ff(*src,"-c:v","libx264","-pix_fmt","yuv420p","-c:a","aac","-shortest",str(OUT/"视频 样本.mp4"))
ff(*src,"-c:v","libx265","-c:a","libopus","-shortest",str(OUT/"视频.mkv"))
ff(*src,"-c:v","libx264","-pix_fmt","yuv420p","-c:a","aac","-shortest",str(OUT/"视频.mov"))
ff(*src,"-c:v","mpeg4","-c:a","libmp3lame","-shortest",str(OUT/"视频.avi"))
ff(*src,"-c:v","libvpx-vp9","-b:v","500k","-c:a","libopus","-shortest",str(OUT/"视频.webm"))
ff(*src,"-c:v","flv","-c:a","libmp3lame","-ar","44100","-shortest",str(OUT/"视频.flv"))
ff(*src,"-c:v","wmv2","-c:a","wmav2","-shortest",str(OUT/"视频.wmv"))
ff(*src,"-c:v","libx264","-pix_fmt","yuv420p","-c:a","aac","-shortest",str(OUT/"视频.m4v"))
ff(*src,"-c:v","libx264","-pix_fmt","yuv420p","-c:a","aac","-shortest","-f","mpegts",str(OUT/"视频.ts"))
ff("-f","lavfi","-i","testsrc2=duration=5:size=352x288:rate=15","-f","lavfi","-i","sine=duration=5","-c:v","h263","-c:a","aac","-ar","16000","-ac","1","-shortest",str(OUT/"视频.3gp"))
ff(*src,"-c:v","mpeg2video","-c:a","mp2","-shortest",str(OUT/"视频.mpg"))
ff(*src,"-c:v","mpeg2video","-c:a","mp2","-shortest","-f","mpeg",str(OUT/"视频.mpeg"))
# ---- 音频（4 秒）
a = ["-f","lavfi","-i","sine=frequency=523:duration=4"]
ff(*a,"-c:a","libmp3lame",str(OUT/"音频 样本.mp3")); ff(*a,str(OUT/"音频.wav")); ff(*a,str(OUT/"音频.flac"))
ff(*a,"-c:a","aac","-f","adts",str(OUT/"音频.aac")); ff(*a,"-c:a","aac",str(OUT/"音频.m4a"))
ff(*a,"-c:a","libvorbis",str(OUT/"音频.ogg")); ff(*a,"-c:a","libopus",str(OUT/"音频.opus"))
ff(*a,"-c:a","wmav2",str(OUT/"音频.wma")); ff(*a,str(OUT/"音频.aiff")); ff(*a,"-f","aiff",str(OUT/"音频.aif"))
# ---- 文档
(OUT/"文本.txt").write_text(f"{MARK}\n第二行 second line\n", "utf-8")
(OUT/"笔记.md").write_text(f"# 标题\n\n{MARK}\n\n- 列表项\n", "utf-8")
(OUT/"表格.csv").write_text(f"名称,数量\n{MARK},3\n苹果,5\n", "utf-8")
tmp = OUT/"_tmp"; tmp.mkdir(exist_ok=True)
def lo(src, fmt, name):
    subprocess.run(["soffice","--headless","--convert-to",fmt,"--outdir",str(tmp),str(src)],check=True,capture_output=True)
    produced = tmp/(Path(src).stem+"."+fmt.split(":")[0]); produced.rename(OUT/name)
for fmt in ("docx","odt","rtf","doc"): lo(OUT/"文本.txt", fmt, f"文档.{fmt}")
lo(OUT/"文本.txt","pdf","文档 样本.pdf")
for fmt in ("xlsx","ods","xls"): lo(OUT/"表格.csv", fmt, f"表格.{fmt}")
from pptx import Presentation
pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = MARK; s.placeholders[1].text = "第二页内容"
pr.slides.add_slide(pr.slide_layouts[1]).shapes.title.text = "第二张"
pr.save(OUT/"演示.pptx")
lo(OUT/"演示.pptx","odp","演示.odp"); lo(OUT/"演示.pptx","ppt","演示.ppt")
# 多页 PDF
import fitz
d = fitz.open()
for i in range(3):
    pg = d.new_page(); pg.insert_text((72,72), f"Page {i+1} MiniBox", fontsize=20)
d.save(OUT/"三页.pdf")
import shutil; shutil.rmtree(tmp)
print("\n".join(sorted(p.name for p in OUT.iterdir())))
